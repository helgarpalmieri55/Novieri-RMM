"""The estate export as a real `.xlsx` workbook, and the audit event that
every format of it owes invariant #4.

WHY A SECOND FORMAT RATHER THAN A REPLACEMENT. `/export/devices.csv` and
`/export/software.csv` are shipped GET routes; this repository cannot see
what outside the console reads them (a scheduled fetch, a CMDB import, a
spreadsheet's own web query), and deleting a route to change a file format
is a breaking change made on behalf of consumers nobody enumerated. The
console's button is a console-side choice and moves to `.xlsx`; the CSV
contract stays.

WHY THE FORMULA GUARD MATTERS MORE HERE, NOT LESS. `app.core.csv_export`
defends a CSV against a spreadsheet's *import parser* deciding a cell that
begins `=`/`+`/`-`/`@` is a formula. A workbook skips that step: openpyxl
writes a string beginning with `=` as `t="f"` -- a real formula in a real
sheet, evaluated the instant the file opens, with no import dialog in
between. The value is agent-reported free text, so a compromised endpoint
naming itself `=HYPERLINK(...)` is the whole attack. These tests pin that
no cell in either workbook is ever a formula cell.
"""

import csv
import io
import zipfile
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.tenancy import org_session
from app.models import AuditEvent, Device
from app.modules.inventory.ingest import ingest_inventory
from app.modules.inventory.schema import InventoryDocument
from tests.test_inventory_api import _seed_inventory
from tests.test_inventory_schema import FULL_DOC

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# The document a compromised or mischievous endpoint reports about itself.
POISONED = {
    "hostname": "=cmd|' /C calc'!A1",
    "serial": "-2+3+cmd|/c calc",
    "software": [{"name": "@SUM(1,1)", "version": "+1", "publisher": "-Acme"}],
}


def _sheet(body: bytes):
    return load_workbook(io.BytesIO(body)).active


def _rows(body: bytes) -> list[list[str]]:
    return [[c.value for c in row] for row in _sheet(body).iter_rows()]


async def _poison(api_env, org_id: UUID, device_id: UUID) -> None:
    async with org_session(org_id, engine=api_env.app_engine) as session:
        device = (
            await session.execute(select(Device).where(Device.id == device_id))
        ).scalar_one()
        await ingest_inventory(
            session, device, InventoryDocument(**{**FULL_DOC, **POISONED})
        )
        await session.commit()


async def test_devices_xlsx_is_a_workbook_with_the_csv_route_s_columns(api_env) -> None:
    org_id = await api_env.create_org("ixlsx")
    await _seed_inventory(api_env, org_id, "xlsx-host")
    reader = await api_env.create_user(org_id, "read-only")
    headers = await api_env.headers(org_id, reader)

    response = await api_env.client.get("/api/v1/export/devices.xlsx", headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX_MEDIA)
    assert 'filename="devices.xlsx"' in response.headers["content-disposition"]

    rows = _rows(response.content)
    # THE SAME NINE COLUMNS AS THE CSV ROUTE, in the same order -- `source` and
    # `provenance` appended 2026-08-26 so a swept estate's workbook can say
    # which rows a person vouched for. Both formats read one row builder, which
    # is what keeps this list and the CSV's from drifting.
    assert rows[0] == [
        "hostname",
        "os_family",
        "os_version",
        "architecture",
        "serial",
        "agent_status",
        "last_seen_at",
        "source",
        "provenance",
    ]
    assert any(r[0] == "xlsx-host" for r in rows[1:])


async def test_software_xlsx_is_a_workbook_with_the_csv_route_s_columns(api_env) -> None:
    org_id = await api_env.create_org("ixlsx")
    await _seed_inventory(api_env, org_id, "xlsx-host")
    reader = await api_env.create_user(org_id, "read-only")
    headers = await api_env.headers(org_id, reader)

    response = await api_env.client.get("/api/v1/export/software.xlsx", headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX_MEDIA)
    assert 'filename="software.xlsx"' in response.headers["content-disposition"]

    rows = _rows(response.content)
    assert rows[0] == ["hostname", "name", "version", "publisher"]
    assert any(r[:2] == ["xlsx-host", "7-Zip"] for r in rows[1:])


async def test_xlsx_and_csv_export_the_same_estate(api_env) -> None:
    """ONE ROW BUILDER, TWO RENDERERS. A second format is only a format if it
    carries the same facts; two independent queries would drift the first
    time a column was added to one of them."""
    org_id = await api_env.create_org("ixlsx")
    await _seed_inventory(api_env, org_id, "same-host")
    reader = await api_env.create_user(org_id, "read-only")
    headers = await api_env.headers(org_id, reader)

    for name, cols in (("devices", 9), ("software", 4)):
        as_csv = await api_env.client.get(f"/api/v1/export/{name}.csv", headers=headers)
        as_xlsx = await api_env.client.get(f"/api/v1/export/{name}.xlsx", headers=headers)
        csv_rows = [r for r in csv.reader(io.StringIO(as_csv.text)) if r]
        xlsx_rows = [["" if c is None else c for c in r] for r in _rows(as_xlsx.content)]
        assert len(csv_rows[0]) == cols
        assert xlsx_rows == csv_rows, name


async def test_xlsx_exports_require_devices_read(api_env) -> None:
    org_id = await api_env.create_org("ixlsx")
    nobody = await api_env.create_user(org_id)
    headers = await api_env.headers(org_id, nobody)
    for path in ("/api/v1/export/devices.xlsx", "/api/v1/export/software.xlsx"):
        assert (await api_env.client.get(path, headers=headers)).status_code == 403, path


async def test_xlsx_exports_are_tenant_scoped(api_env) -> None:
    """Invariant #5. A whole-estate export is the single worst place for a
    missing `org_id` filter, because one file carries every row it leaked."""
    org_a = await api_env.create_org("ixlsx-a")
    org_b = await api_env.create_org("ixlsx-b")
    await _seed_inventory(api_env, org_a, "mine-host")
    await _seed_inventory(api_env, org_b, "theirs-host")
    reader = await api_env.create_user(org_a, "read-only")
    headers = await api_env.headers(org_a, reader)

    devices = await api_env.client.get("/api/v1/export/devices.xlsx", headers=headers)
    flat = {c for row in _rows(devices.content) for c in row}
    assert "mine-host" in flat
    assert "theirs-host" not in flat

    software = await api_env.client.get("/api/v1/export/software.xlsx", headers=headers)
    flat = {c for row in _rows(software.content) for c in row}
    assert "mine-host" in flat
    assert "theirs-host" not in flat


async def test_no_cell_in_either_workbook_is_a_formula(api_env) -> None:
    """THE ONE THAT MATTERS. `escape_formula` neutralises a CSV cell for a
    spreadsheet's import parser; a workbook has no import parser, so the
    equivalent defence is that the cell is never written as a formula in the
    first place. Asserted over EVERY cell, not only the free-text columns, so
    a column added later without thinking about it is still inert."""
    org_id = await api_env.create_org("ixlsx")
    device_id = await _seed_inventory(api_env, org_id)
    await _poison(api_env, org_id, device_id)
    reader = await api_env.create_user(org_id, "read-only")
    headers = await api_env.headers(org_id, reader)

    for path in ("/api/v1/export/devices.xlsx", "/api/v1/export/software.xlsx"):
        response = await api_env.client.get(path, headers=headers)
        assert response.status_code == 200, response.text

        # THE FALSIFIER, READ OFF THE FILE ITSELF. `t="f"` on a `<c>`, or an
        # `<f>` child of one, is what makes a cell a formula in the
        # SpreadsheetML a spreadsheet actually opens. Asserted against the
        # bytes rather than against openpyxl's model, so this stays true of
        # the artefact even if the library's inference changes.
        sheet_xml = zipfile.ZipFile(io.BytesIO(response.content)).read(
            "xl/worksheets/sheet1.xml"
        ).decode()
        assert 't="f"' not in sheet_xml, path
        assert "<f>" not in sheet_xml and "<f " not in sheet_xml, path

        # And through the model, over EVERY cell rather than only the
        # free-text columns -- a column added later without reading
        # `_write_text_cell` is still inert. (An empty cell round-trips as
        # `inlineStr`/None rather than `s`; both are string types, neither is
        # a formula.)
        sheet = _sheet(response.content)
        types = {cell.data_type for row in sheet.iter_rows() for cell in row}
        assert types <= {"s", "inlineStr"}, f"{path} wrote a non-string cell: {types}"


async def test_poisoned_free_text_keeps_its_text_and_is_marked_forced_text(api_env) -> None:
    """`quotePrefix` is Excel's own "this cell is text, not a formula" marker
    -- what a leading `'` MEANS to a CSV import, expressed natively. So the
    operator reads the endpoint's real reported name rather than a mangled
    one, and the columns the CSV route leaves alone (a validated enum, a
    server-formatted timestamp) are left alone here too."""
    org_id = await api_env.create_org("ixlsx")
    device_id = await _seed_inventory(api_env, org_id)
    await _poison(api_env, org_id, device_id)
    reader = await api_env.create_user(org_id, "read-only")
    headers = await api_env.headers(org_id, reader)

    response = await api_env.client.get("/api/v1/export/devices.xlsx", headers=headers)
    sheet = _sheet(response.content)
    header = [c.value for c in next(sheet.iter_rows())]
    row = next(
        r for r in sheet.iter_rows(min_row=2) if r[header.index("serial")].value.endswith("calc")
    )
    hostname = row[header.index("hostname")]
    assert hostname.value == POISONED["hostname"]
    assert hostname.quotePrefix is True
    assert row[header.index("serial")].quotePrefix is True
    # A pattern-validated enum can never begin with a dangerous character, so
    # it is not marked -- the same column-by-column judgement the CSV route
    # documents, not a blanket sweep.
    os_family = row[header.index("os_family")]
    assert os_family.value == "windows"
    assert os_family.quotePrefix is not True

    response = await api_env.client.get("/api/v1/export/software.xlsx", headers=headers)
    sheet = _sheet(response.content)
    sheader = [c.value for c in next(sheet.iter_rows())]
    srow = next(
        r for r in sheet.iter_rows(min_row=2) if r[sheader.index("name")].value == "@SUM(1,1)"
    )
    assert [srow[sheader.index(c)].quotePrefix for c in ("name", "version", "publisher")] == [
        True,
        True,
        True,
    ]


async def test_every_estate_export_writes_one_audit_event(api_env) -> None:
    """INVARIANT #4. A whole-tenant inventory export is a privileged read of
    the entire estate leaving the platform as a file -- the same class of act
    as `report.exported`, `patch.report_exported` and `bcp.report_exported`,
    each of which cites this invariant. The CSV route had shipped writing
    nothing at all, so a fleet's hardware, serials and installed software
    could be walked out with no record that it happened."""
    org_id = await api_env.create_org("ixlsx")
    await _seed_inventory(api_env, org_id, "audited-host")
    reader = await api_env.create_user(org_id, "read-only")
    headers = await api_env.headers(org_id, reader)

    for path in (
        "/api/v1/export/devices.csv",
        "/api/v1/export/devices.xlsx",
        "/api/v1/export/software.csv",
        "/api/v1/export/software.xlsx",
    ):
        assert (await api_env.client.get(path, headers=headers)).status_code == 200, path

    async with org_session(org_id, engine=api_env.app_engine) as session:
        events = [
            e
            for e in (
                await session.execute(
                    select(AuditEvent)
                    .where(AuditEvent.org_id == org_id, AuditEvent.action == "inventory.exported")
                    .order_by(AuditEvent.id)
                )
            ).scalars()
        ]

    assert len(events) == 4
    assert [(e.meta["dataset"], e.meta["format"]) for e in events] == [
        ("devices", "csv"),
        ("devices", "xlsx"),
        ("software", "csv"),
        ("software", "xlsx"),
    ]
    # The COUNT, never the rows. `audit_events` is append-only and unprunable
    # by `app_user`, so a row of exported estate written here could not be
    # taken back.
    assert all(isinstance(e.meta["row_count"], int) for e in events)
    assert all(e.actor_type == "user" and e.actor_id is not None for e in events)
    assert all(e.target_type == "organization" for e in events)
